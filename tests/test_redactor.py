# basic tests for the redaction pipeline
import unittest
import zipfile
from pathlib import Path
from tempfile import TemporaryDirectory
from types import SimpleNamespace

from docx import Document

from pii_redactor.recognizers import ConflictResolver, PiiRecognizer
from pii_redactor.replacement_generator import ReplacementGenerator
from pii_redactor.docx_redaction_service import DocxRedactor, DocxRedactionService
from pii_redactor.indexer import DocxTextIndexer
from pii_redactor.models import EntitySpan, Replacement, RunSegment, TextUnit
from pii_redactor.span_propagator import SpanPropagator
from pii_redactor.docx_field_sanitizer import DocxFieldSanitizer


class FakeRun:
    def __init__(self, text):
        self.text = text


# --- detection stuff ---
class DetectionTests(unittest.TestCase):
    def setUp(self):
        self.detector = PiiRecognizer(use_optional_ner=False)

    def pairs(self, text):
        return {(span.entity_type, " ".join(span.text.split())) for span in self.detector.detect_text(text)}

    # check that all the basic PII types get picked up
    def test_structured_pii_detection(self):
        text = (
            "DOB: 03/09/1992 SSN: 123-45-6789 Card: 4111 1111 1111 1111 "
            "IP 192.168.1.10 Email: rashi.patil@gmail.com Tel: +91 22 6807 7100"
        )
        pairs = self.pairs(text)
        self.assertIn(("dob", "03/09/1992"), pairs)
        self.assertIn(("ssn", "123-45-6789"), pairs)
        self.assertIn(("credit_card", "4111 1111 1111 1111"), pairs)
        self.assertIn(("ip_address", "192.168.1.10"), pairs)
        self.assertIn(("email", "rashi.patil@gmail.com"), pairs)
        self.assertIn(("phone", "+91 22 6807 7100"), pairs)

    # docx sometimes splits URLs with spaces, gotta handle that
    def test_spaced_docx_url_detection(self):
        pairs = self.pairs("Website: www.kshinternational. com")
        self.assertIn(("url", "www.kshinternational. com"), pairs)

    def test_names_companies_and_addresses(self):
        text = (
            "Nuvama Wealth Management Limited, Contact Person: Lokesh Shah. "
            "Registered Office: 11/3, Village Birdewadi, Chakan Taluka - Khed, "
            "Pune - 410 501, Maharashtra, India;"
        )
        pairs = self.pairs(text)
        self.assertIn(("company", "Nuvama Wealth Management Limited"), pairs)
        self.assertIn(("person", "Lokesh Shah"), pairs)
        self.assertTrue(any(pair[0] == "address" and "410 501" in pair[1] for pair in pairs))

    # don't want random dates being tagged as DOB
    def test_offer_date_is_not_dob(self):
        pairs = self.pairs("Bid closes on December 18, 2025.")
        self.assertNotIn(("dob", "December 18, 2025"), pairs)


class ReplacementTests(unittest.TestCase):
    # make sure same name always maps to same fake
    def test_same_original_maps_to_same_fake_value(self):
        generator = ReplacementGenerator()
        first = generator.replacement_for("person", "Sarthak Malvadkar").fake_value
        second = generator.replacement_for("person", "Sarthak   Malvadkar").fake_value
        self.assertEqual(first, second)


class ConflictResolverTests(unittest.TestCase):
    # overlap resolution test
    def test_higher_priority_span_wins_overlap(self):
        spans = [
            EntitySpan("company", "example.com", 10, 21, 0.8, "company"),
            EntitySpan("email", "user@example.com", 5, 21, 0.99, "email"),
        ]
        resolved = ConflictResolver.resolve(spans)
        self.assertEqual(len(resolved), 1)
        self.assertEqual(resolved[0].entity_type, "email")


class PropagationTests(unittest.TestCase):
    # if we already know "Kushal" is a person, it should get caught even in ALL CAPS
    def test_propagates_known_person_case_insensitively(self):
        unit = TextUnit(
            "u1",
            "body",
            "OFFERED SHARES BY KUSHAL SUBBAYYA HEGDE AGGREGATING UP TO Rs. 100.",
            [],
        )
        lookup = {
            ("person", "kushal subbayya hegde"): Replacement(
                "Kushal Subbayya Hegde",
                "person",
                "Aarav Menon",
            )
        }
        spans_by_unit = {"u1": []}

        SpanPropagator().add_propagation([unit], spans_by_unit, lookup)

        pairs = {(span.entity_type, span.text) for span in spans_by_unit["u1"]}
        self.assertIn(("person", "KUSHAL SUBBAYYA HEGDE"), pairs)


class RunRedactionTests(unittest.TestCase):
    # names can get split across multiple docx runs, this checks we handle that
    def test_redacts_across_multiple_runs(self):
        runs = [FakeRun("Contact Person: "), FakeRun("Sarthak"), FakeRun(" "), FakeRun("Malvadkar")]
        text = "".join(run.text for run in runs)
        segments = []
        offset = 0
        for run in runs:
            segments.append(RunSegment(run=run, start=offset, end=offset + len(run.text)))
            offset += len(run.text)
        unit = TextUnit("u1", "body", text, segments)
        span = EntitySpan("person", "Sarthak Malvadkar", text.index("Sarthak"), len(text), 0.9, "test")
        replacement = SimpleNamespace(fake_value="Aarav Mehta")

        service = DocxRedactor(scanner=PiiRecognizer(use_optional_ner=False))
        service._apply_spans(unit, [span], {span: replacement})

        self.assertEqual("".join(run.text for run in runs), "Contact Person: Aarav Mehta")


class DocxIndexerTests(unittest.TestCase):
    # merged cells in tables were causing duplicates earlier, this should be fixed now
    def test_merged_table_cell_is_indexed_once(self):
        doc = Document()
        table = doc.add_table(rows=1, cols=2)
        merged = table.cell(0, 0).merge(table.cell(0, 1))
        merged.text = "Contact email user@example.com"

        units = DocxTextIndexer().index(doc)
        matching_units = [unit for unit in units if "user@example.com" in unit.text]

        self.assertEqual(len(matching_units), 1)


# FIXME: might need more tests for other field types later
class DocxFieldSanitizerTests(unittest.TestCase):
    # hyperlinks in docx have the email in the XML instrText, gotta sanitize that too
    def test_removes_original_email_from_hyperlink_field_instruction(self):
        with TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "sample.docx"
            with zipfile.ZipFile(path, "w") as archive:
                archive.writestr("[Content_Types].xml", "<Types/>")
                archive.writestr(
                    "word/document.xml",
                    '<w:instrText> HYPERLINK "mailto:person@example.com"</w:instrText>',
                )

            DocxFieldSanitizer().sanitize(
                path,
                [Replacement("person@example.com", "email", "user1234@example.com")],
            )

            with zipfile.ZipFile(path) as archive:
                xml = archive.read("word/document.xml").decode("utf-8")
            self.assertNotIn("person@example.com", xml)
            self.assertIn("user1234@example.com", xml)


if __name__ == "__main__":
    unittest.main()
